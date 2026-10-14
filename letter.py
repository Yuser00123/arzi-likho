"""
letter.py — assembling the final letter.

Fixed skeleton (date, सेवा में block, विषय, salutation, प्रार्थना, भवदीय, signature, संलग्नक,
"कहाँ जमा करें" box) comes from forms.py; the narrative paragraphs come from the draft JSON
(written by Gemma, or the deterministic mock). Output is plain text so the operator can edit
it in the browser before printing, and a .docx for the café's Word workflow.
"""
from __future__ import annotations

import os
import re
import tempfile
from datetime import datetime
from typing import Dict

import forms


def _today() -> str:
    return datetime.now().strftime("%d/%m/%Y")


def _verb_suffix(gender: str) -> str:
    if gender.strip() == "महिला":
        return "रही हूँ"
    if gender.strip() == "पुरुष":
        return "रहा हूँ"
    return "रहा/रही हूँ"


def _intro(fields: Dict[str, str], district: str) -> str:
    return (
        f"सविनय निवेदन है कि मैं, {forms.sv(fields, 'name', 'नाम लिखें')}, "
        f"पिता/पति — {forms.sv(fields, 'father', 'पिता/पति का नाम लिखें')}, "
        f"निवासी {forms.sv(fields, 'address', 'पूरा पता लिखें')}, जिला {district} (उत्तर प्रदेश), "
        f"यह अर्ज़ी आपके समक्ष प्रस्तुत कर {_verb_suffix(fields.get('gender', ''))}।"
    )


def _attachments_md(matter: dict) -> str:
    if not matter.get("attachments"):
        return ""
    lines = ["संलग्नक:"] + [f"{i}. {a}" for i, a in enumerate(matter["attachments"], 1)]
    return "\n".join(lines)


def kahan_box(matter: dict) -> str:
    """The 'where to file / what next' box — pure template, no model."""
    rows = ["### 🏢 कहाँ जमा करें · आगे क्या करें"]
    if matter.get("online"):
        rows.append(f"**ऑनलाइन:** {matter['online']}")
    if matter.get("helpline"):
        rows.append(f"**हेल्पलाइन:** {matter['helpline']}")
    if matter.get("fee"):
        rows.append(f"**शुल्क:** {matter['fee']}")
    if matter.get("time_limit"):
        rows.append(f"**समय-सीमा:** {matter['time_limit']}")
    if matter.get("legal"):
        rows.append(f"**कानूनी आधार:** {matter['legal']}")
    if matter.get("attachments"):
        rows.append("**संलग्नक (साथ रखें):** " + " · ".join(matter["attachments"]))
    rows.append("<small>कार्यालय, शुल्क और पोर्टल समय-समय पर बदलते हैं — जमा करने से पहले कार्यालय/जन सुविधा केंद्र से पुष्टि कर लें।</small>")
    return "\n\n".join(rows)


# Leading connectives the model sometimes prepends to the request — the letter's
# "अतः प्रार्थना है कि ..." prefix already does that job.
_REQ_LEADS = ("अतः", "अत:", "इसलिए", "प्रार्थना है कि", "आपसे विनम्र निवेदन है कि", "विनम्र निवेदन है कि", "निवेदन है कि", "कि")


def _clean_request(r: str) -> str:
    r = (r or "").strip()
    changed = True
    while changed and r:
        changed = False
        for w in _REQ_LEADS:
            if r == w:
                return ""
            if r.startswith(w + " "):
                r = r[len(w):].strip()
                changed = True
    return r


def compose(matter: dict, district: str, fields: Dict[str, str], draft: dict) -> Dict[str, str]:
    """Everything the UI needs: the editable letter text + the boxes around it."""
    if matter.get("style") == "biodata":
        return _compose_biodata(matter, district, fields, draft)
    addressee = forms.addressee_lines(matter, fields, district)
    subject = matter["subject_fn"](fields)
    paras = draft.get("body") or []
    letter = [
        f"दिनांक: {_today()}",
        "",
        "सेवा में,",
        *[a for a in addressee],
        "",
        f"विषय: {subject}",
        "",
        "महोदय,",
        "",
        _intro(fields, district),
        "",
        *"\n\n".join(paras).split("\n"),
        "",
        f"अतः प्रार्थना है कि {_clean_request(draft.get('request') or matter['request'])}",
        "",
        "भवदीय,",
        forms.sv(fields, "name", "नाम लिखें"),
        forms.sv(fields, "father", "पिता/पति का नाम लिखें"),
        f"{forms.sv(fields, 'address', 'पूरा पता लिखें')}, जिला {district}",
        forms.sv(fields, "mobile", "मोबाइल नंबर लिखें"),
    ]
    att = _attachments_md(matter)
    if att:
        letter += ["", *att.split("\n")]
    for cc in matter.get("cc", []):
        letter += ["", cc.format(district=district)]
    english = draft.get("english") or {}
    english_md = "\n".join([
        f"**Subject: {english.get('subject', matter.get('subject_en', ''))}**", "",
        "Sir/Madam,", "",
        *[p for p in english.get("body", [])],
        "",
        english.get("request", ""),
        "",
        "Yours faithfully,",
        forms.sv(fields, "name", "नाम लिखें"),
    ])
    return {
        "letter": "\n".join(letter),
        "kahan": kahan_box(matter),
        "english": english_md.strip(),
        "whatsapp": draft.get("whatsapp") or "",
        "subject": subject,
    }


def _compose_biodata(matter: dict, district: str, fields: Dict[str, str], draft: dict) -> Dict[str, str]:
    def block(title: str, key: str) -> list:
        v = (fields.get(key) or "").strip()
        if not v:
            return []
        lines = [v] if "\n" not in v else v.split("\n")
        return [f"__{title}__", *[f"• {ln.strip()}" for ln in lines if ln.strip()]]

    bd = [
        f"⸻ {forms.sv(fields, 'purpose', 'बायोडाटा')} ⸻", "",
        "__व्यक्तिगत जानकारी__",
        f"नाम: {forms.sv(fields, 'name', 'नाम लिखें')}",
        f"जन्म तिथि: {forms.sv(fields, 'dob', 'जन्म तिथि लिखें')}",
        f"पता: {forms.sv(fields, 'address', 'पूरा पता लिखें')}, जिला {district}",
        f"मोबाइल: {forms.sv(fields, 'mobile', 'मोबाइल नंबर लिखें')}",
    ]
    about = (draft.get("about") or "").strip()
    if about:
        bd += ["", "__परिचय__", about]
    for title, key in (("शिक्षा", "education"), ("कार्य / अनुभव", "work"), ("परिवार", "family"), ("अन्य जानकारी", "extra")):
        b = block(title, key)
        if b:
            bd += ["", *b]
    english_about = (draft.get("english_about") or "").strip()
    english_md = f"**About**\n\n{english_about}" if english_about else ""
    return {
        "letter": "\n".join(bd),
        "kahan": "",
        "english": english_md,
        "whatsapp": draft.get("whatsapp") or "",
        "subject": matter["subject_fn"](fields),
    }


# ----------------------------------------------------------------------------- .docx
def build_docx(text: str, out_path: str | None = None) -> str:
    """The letter's plain text → an A4 Word file (Devanagari-safe). Cafés live in Word."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(22)
    sec.top_margin = sec.bottom_margin = Mm(20)

    style = doc.styles["Normal"]
    style.font.name = "Nirmala UI"          # Windows ships Nirmala UI/Mangal for Devanagari
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Nirmala UI")
    rfonts.set(qn("w:hAnsi"), "Nirmala UI")
    rfonts.set(qn("w:cs"), "Nirmala UI")

    for raw in text.split("\n"):
        line = raw.rstrip()
        p = doc.add_paragraph()
        run = p.add_run(line)
        s = line.strip()
        _BD_HEADS = ("व्यक्तिगत जानकारी", "शिक्षा", "कार्य / अनुभव", "परिवार", "अन्य जानकारी", "परिचय")
        if s.startswith("विषय:") or s.startswith("Subject:") or s.startswith("संलग्नक") or s in _BD_HEADS or (s.startswith("⸻") and s.endswith("⸻")):
            run.bold = True
        if s.startswith("⸻") and s.endswith("⸻"):
            p.alignment = 1  # centre the bio-data heading
    path = out_path or os.path.join(tempfile.gettempdir(), "arzi-likho.docx")
    doc.save(path)
    return path


def docx_from_letter(letter_text: str, matter_label: str) -> str:
    safe = re.sub(r"[^\w-]+", "_", matter_label).strip("_")[:40] or "arzi"
    fname = os.path.join(tempfile.gettempdir(), f"arzi-{safe}-{datetime.now().strftime('%Y%m%d-%H%M%S')}.docx")
    return build_docx(letter_text, fname)
